"""Convert an Excel workbook to plain text so that `git diff` can show cell-level changes.

Use it as a git `textconv` driver. See the "Diff Excel files with git" section of `backend/README.md`.
"""
import sys
from datetime import date, time
from typing import Iterable, Iterator

from openpyxl import load_workbook
from openpyxl.workbook.defined_name import DefinedName


def _format_name(name: DefinedName) -> str:
    """Format a defined name as `<name>: <reference>` with its optional flags.

    Args:
        name: The defined name (named range) to format.

    Returns:
        The text for the name, without any scope prefix.
    """
    text = f"{name.name}: {name.attr_text}"
    if name.hidden:
        text += " [hidden]"
    if name.comment:
        text += f" comment={name.comment!r}"
    return text


def _sorted_names(names: Iterable[DefinedName]) -> list[DefinedName]:
    """Sort defined names alphabetically so that their internal order never shows as a change."""
    return sorted(names, key=lambda n: n.name)


def dump_workbook(path: str) -> Iterator[str]:
    """Yield one text line per item of the workbook: cells, merges, validations and named ranges.

    Formulas appear as formula text. Values use `repr()` so that `"1"` and `1` differ.

    Args:
        path: Path to the `.xlsx` or `.xlsm` file.

    Yields:
        The lines of the text representation, in workbook order.
    """
    wb = load_workbook(path, data_only=False)
    for ws in wb.worksheets:
        yield f"=== Sheet: {ws.title} [state={ws.sheet_state}] ==="
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    value = cell.value.isoformat() if isinstance(cell.value, (date, time)) else repr(cell.value)
                    yield f"{ws.title}!{cell.coordinate}: {value}"
        for merged in sorted(str(r) for r in ws.merged_cells.ranges):
            yield f"{ws.title} merged: {merged}"
        for dv in ws.data_validations.dataValidation:
            yield f"{ws.title} validation {dv.sqref}: type={dv.type} formula1={dv.formula1}"
        for name in _sorted_names(ws.defined_names.values()):
            yield f"{ws.title} name {_format_name(name)}"
    for name in _sorted_names(wb.defined_names.values()):
        yield f"name {_format_name(name)}"


def main() -> int:
    """Print the text representation of the workbook given as the only argument.

    Returns:
        0 on success, 1 on a bad argument or an unreadable file.
    """
    if len(sys.argv) != 2:
        print("usage: xlsx_textconv.py <workbook.xlsx>", file=sys.stderr)
        return 1
    try:
        for line in dump_workbook(sys.argv[1]):
            print(line)
    except Exception as e:
        print(f"xlsx_textconv: cannot read {sys.argv[1]}: {e}", file=sys.stderr)
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
