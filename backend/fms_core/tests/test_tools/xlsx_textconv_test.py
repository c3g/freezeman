import os
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.workbook.defined_name import DefinedName

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "tools"))
from xlsx_textconv import dump_workbook  # noqa: E402

VALID_TEMPLATES = Path(__file__).resolve().parents[1] / "valid_templates"


def _make_workbook(path: Path, cell: str = "a", global_ref: str = "Data!$A$1:$A$5", local_ref: str = "Data!$B$1") -> str:
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws["A1"] = "header"
    ws["A2"] = cell
    wb.defined_names["GlobalRange"] = DefinedName("GlobalRange", attr_text=global_ref)
    ws.defined_names["LocalRange"] = DefinedName("LocalRange", attr_text=local_ref)
    wb.save(path)
    return str(path)


def test_dump_real_template():
    path = next(VALID_TEMPLATES.glob("*.xlsx"))
    lines = list(dump_workbook(str(path)))
    assert any(line.startswith("=== Sheet: ") for line in lines)
    assert any("!A1:" in line or "!B1:" in line for line in lines) or len(lines) > 1


def test_single_cell_change(tmp_path):
    old = list(dump_workbook(_make_workbook(tmp_path / "old.xlsx", cell="a")))
    new = list(dump_workbook(_make_workbook(tmp_path / "new.xlsx", cell="b")))
    assert [l for l in old if l not in new] == ["Data!A2: 'a'"]
    assert [l for l in new if l not in old] == ["Data!A2: 'b'"]


def test_named_range_change(tmp_path):
    old = list(dump_workbook(_make_workbook(tmp_path / "old.xlsx")))
    new = list(dump_workbook(_make_workbook(tmp_path / "new.xlsx", global_ref="Data!$A$1:$A$9")))
    assert "name GlobalRange: Data!$A$1:$A$5" in old
    assert [l for l in old if l not in new] == ["name GlobalRange: Data!$A$1:$A$5"]
    assert [l for l in new if l not in old] == ["name GlobalRange: Data!$A$1:$A$9"]

    newer = list(dump_workbook(_make_workbook(tmp_path / "newer.xlsx", local_ref="Data!$C$1")))
    assert [l for l in old if l not in newer] == ["Data name LocalRange: Data!$B$1"]
    assert [l for l in newer if l not in old] == ["Data name LocalRange: Data!$C$1"]


def test_by_column_order(tmp_path):
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Data"
    ws["A1"], ws["B1"], ws["A2"], ws["B2"] = "a1", "b1", "a2", "b2"
    path = tmp_path / "order.xlsx"
    wb.save(path)

    def cells(by_column: bool) -> list[str]:
        return [l.split(":")[0] for l in dump_workbook(str(path), by_column=by_column) if l.startswith("Data!")]

    assert cells(False) == ["Data!A1", "Data!B1", "Data!A2", "Data!B2"]
    assert cells(True) == ["Data!A1", "Data!A2", "Data!B1", "Data!B2"]
